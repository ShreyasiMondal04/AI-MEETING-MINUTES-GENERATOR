import json
from pathlib import Path
from typing import Dict, Any, Union
from datetime import datetime

from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether, HRFlowable
)
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT


def export_pdf(minutes: Dict[str, Any], output_path: Union[str, Path]) -> Path:
    """
    Generate an executive-ready, beautifully styled PDF meeting minutes document
    using reportlab.
    """
    out_file = Path(output_path)
    out_file.parent.mkdir(parents=True, exist_ok=True)

    doc = SimpleDocTemplate(
        str(out_file),
        pagesize=letter,
        leftMargin=40,
        rightMargin=40,
        topMargin=40,
        bottomMargin=40
    )

    styles = getSampleStyleSheet()

    # Custom styles
    primary_color = colors.HexColor("#1e3a8a")  # Deep Navy Blue
    accent_color = colors.HexColor("#4338ca")   # Indigo
    dark_gray = colors.HexColor("#1f2937")
    light_bg = colors.HexColor("#f8fafc")
    border_color = colors.HexColor("#cbd5e1")

    title_style = ParagraphStyle(
        "DocTitle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=20,
        leading=24,
        textColor=primary_color,
        spaceAfter=6
    )

    subtitle_style = ParagraphStyle(
        "DocSubtitle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#64748b"),
        spaceAfter=15
    )

    h1_style = ParagraphStyle(
        "SectionHeading",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=13,
        leading=17,
        textColor=primary_color,
        spaceBefore=12,
        spaceAfter=6
    )

    body_style = ParagraphStyle(
        "BodyTextCustom",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9.5,
        leading=14,
        textColor=dark_gray,
        spaceAfter=6
    )

    bullet_style = ParagraphStyle(
        "BulletCustom",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=13,
        textColor=dark_gray,
        leftIndent=15,
        spaceAfter=4
    )

    callout_style = ParagraphStyle(
        "CalloutText",
        parent=styles["Normal"],
        fontName="Helvetica-Oblique",
        fontSize=9.5,
        leading=14,
        textColor=colors.HexColor("#0f172a")
    )

    table_header_style = ParagraphStyle(
        "TableHeader",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=9,
        leading=12,
        textColor=colors.white
    )

    table_cell_style = ParagraphStyle(
        "TableCell",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8.5,
        leading=12,
        textColor=dark_gray
    )

    story = []

    # Title Banner
    title_text = minutes.get("title", "AI Meeting Minutes Report")
    story.append(Paragraph(title_text, title_style))
    date_str = minutes.get("date", datetime.now().strftime("%B %d, %Y"))
    duration = minutes.get("duration_formatted", "N/A")
    attendees = ", ".join(minutes.get("participants", ["Unknown Attendees"]))
    story.append(Paragraph(f"Date: <b>{date_str}</b> | Duration: <b>{duration}</b> | Attendees: <b>{attendees}</b>", subtitle_style))
    story.append(HRFlowable(width="100%", thickness=1.5, color=primary_color, spaceAfter=12))

    # Executive Summary Card
    story.append(Paragraph("Executive Summary", h1_style))
    summary_text = minutes.get("summary", "No summary available.")
    summary_table = Table(
        [[Paragraph(summary_text, callout_style)]],
        colWidths=[530]
    )
    summary_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), light_bg),
        ('BOX', (0, 0), (-1, -1), 1, border_color),
        ('LINEBEFORE', (0, 0), (0, -1), 3.5, accent_color),
        ('TOPPADDING', (0, 0), (-1, -1), 8),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
        ('LEFTPADDING', (0, 0), (-1, -1), 12),
        ('RIGHTPADDING', (0, 0), (-1, -1), 12),
    ]))
    story.append(summary_table)
    story.append(Spacer(1, 10))

    # Key Discussion Points
    key_points = minutes.get("key_points", [])
    if key_points:
        story.append(Paragraph("Key Discussion Points", h1_style))
        for pt in key_points:
            story.append(Paragraph(f"• {pt}", bullet_style))
        story.append(Spacer(1, 10))

    # Decisions Made
    decisions = minutes.get("decisions", [])
    if decisions:
        story.append(Paragraph("Decisions Made", h1_style))
        for dec in decisions:
            story.append(Paragraph(f"✔ <b>{dec}</b>", bullet_style))
        story.append(Spacer(1, 10))

    # Action Items Table
    action_items = minutes.get("action_items", [])
    if action_items:
        story.append(Paragraph("Action Items & Deliverables", h1_style))
        headers = [
            Paragraph("Task Description", table_header_style),
            Paragraph("Assignee", table_header_style),
            Paragraph("Deadline", table_header_style),
            Paragraph("Priority", table_header_style)
        ]
        table_rows = [headers]
        for item in action_items:
            task = item.get("task", "")
            owner = item.get("responsible_person", "Not identified")
            deadline = item.get("deadline", "Not specified")
            priority = item.get("priority", "Normal")
            table_rows.append([
                Paragraph(task, table_cell_style),
                Paragraph(f"<b>{owner}</b>", table_cell_style),
                Paragraph(deadline, table_cell_style),
                Paragraph(priority, table_cell_style)
            ])

        action_table = Table(table_rows, colWidths=[270, 95, 95, 70])
        action_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), primary_color),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('GRID', (0, 0), (-1, -1), 0.5, border_color),
            ('TOPPADDING', (0, 0), (-1, -1), 5),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, light_bg]),
        ]))
        story.append(action_table)
        story.append(Spacer(1, 12))

    # Speaker Analytics & Talk Time
    speaker_analytics = minutes.get("speaker_analytics", {})
    if speaker_analytics and speaker_analytics.get("speakers"):
        story.append(Paragraph("Speaker Participation & Analytics", h1_style))
        spk_headers = [
            Paragraph("Speaker", table_header_style),
            Paragraph("Speaking Time", table_header_style),
            Paragraph("Share %", table_header_style),
            Paragraph("Turns", table_header_style),
            Paragraph("Word Count", table_header_style)
        ]
        spk_rows = [spk_headers]
        for s in speaker_analytics["speakers"]:
            spk_rows.append([
                Paragraph(f"<b>{s['speaker']}</b>", table_cell_style),
                Paragraph(f"{s['talk_time_sec']}s", table_cell_style),
                Paragraph(f"{s['share_percentage']}%", table_cell_style),
                Paragraph(str(s['turns']), table_cell_style),
                Paragraph(str(s['word_count']), table_cell_style)
            ])
        spk_table = Table(spk_rows, colWidths=[130, 100, 100, 100, 100])
        spk_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#334155")),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('GRID', (0, 0), (-1, -1), 0.5, border_color),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, light_bg]),
        ]))
        story.append(spk_table)
        story.append(Spacer(1, 15))

    # Speaker Transcript Appendix
    transcript_segments = minutes.get("speaker_contributions", [])
    if transcript_segments:
        story.append(Paragraph("Transcript Appendix", h1_style))
        for entry in transcript_segments:
            spk = entry.get("speaker", "Speaker")
            txt = entry.get("text", "")
            ts = entry.get("timestamp", "")
            line = f"<b>{spk}</b> {f'<i>{ts}</i> ' if ts else ''}: {txt}"
            story.append(Paragraph(line, bullet_style))

    doc.build(story)
    return out_file


def export_docx(minutes: Dict[str, Any], output_path: Union[str, Path]) -> Path:
    """
    Generate an editable, formatted Word (.docx) meeting minutes document.
    """
    out_file = Path(output_path)
    out_file.parent.mkdir(parents=True, exist_ok=True)

    doc = docx.Document()

    # Document Title
    title = doc.add_heading(minutes.get("title", "Meeting Minutes"), level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_run = title.runs[0]
    title_run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    # Subtitle / Metadata
    meta_p = doc.add_paragraph()
    meta_p.paragraph_format.space_after = Pt(14)
    meta_p.add_run(f"Date: {minutes.get('date', datetime.now().strftime('%B %d, %Y'))} | ").bold = True
    meta_p.add_run(f"Duration: {minutes.get('duration_formatted', 'N/A')} | ").bold = True
    attendees = ", ".join(minutes.get("participants", ["Unknown Attendees"]))
    meta_p.add_run(f"Attendees: {attendees}").bold = True

    # Executive Summary
    doc.add_heading("1. Executive Summary", level=1)
    summary_p = doc.add_paragraph(minutes.get("summary", "No summary generated."))
    summary_p.paragraph_format.space_after = Pt(12)

    # Key Points
    key_points = minutes.get("key_points", [])
    if key_points:
        doc.add_heading("2. Key Discussion Points", level=1)
        for pt in key_points:
            doc.add_paragraph(pt, style="List Bullet")

    # Decisions Made
    decisions = minutes.get("decisions", [])
    if decisions:
        doc.add_heading("3. Decisions Made", level=1)
        for dec in decisions:
            p = doc.add_paragraph(dec, style="List Bullet")
            p.runs[0].bold = True

    # Action Items Table
    action_items = minutes.get("action_items", [])
    if action_items:
        doc.add_heading("4. Action Items & Next Steps", level=1)
        table = doc.add_table(rows=1, cols=4)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = 'Light Shading Accent 1'

        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = "Task"
        hdr_cells[1].text = "Responsible Person"
        hdr_cells[2].text = "Deadline"
        hdr_cells[3].text = "Priority"

        for item in action_items:
            row_cells = table.add_row().cells
            row_cells[0].text = str(item.get("task", ""))
            row_cells[1].text = str(item.get("responsible_person", "Not identified"))
            row_cells[2].text = str(item.get("deadline", "Not specified"))
            row_cells[3].text = str(item.get("priority", "Normal"))

    # Speaker Analytics
    analytics = minutes.get("speaker_analytics", {})
    if analytics and analytics.get("speakers"):
        doc.add_heading("5. Speaker Participation Metrics", level=1)
        table = doc.add_table(rows=1, cols=5)
        table.style = 'Light Shading Accent 1'
        hdr = table.rows[0].cells
        hdr[0].text = "Speaker"
        hdr[1].text = "Speaking Time"
        hdr[2].text = "Share %"
        hdr[3].text = "Turns"
        hdr[4].text = "Words"

        for s in analytics["speakers"]:
            r = table.add_row().cells
            r[0].text = s["speaker"]
            r[1].text = f"{s['talk_time_sec']}s"
            r[2].text = f"{s['share_percentage']}%"
            r[3].text = str(s["turns"])
            r[4].text = str(s["word_count"])

    # Full Transcript
    transcript = minutes.get("speaker_contributions", [])
    if transcript:
        doc.add_heading("6. Transcript Appendix", level=1)
        for entry in transcript:
            p = doc.add_paragraph()
            p.add_run(f"{entry.get('speaker', 'Speaker')}: ").bold = True
            p.add_run(entry.get("text", ""))

    doc.save(str(out_file))
    return out_file


def export_markdown(minutes: Dict[str, Any], output_path: Union[str, Path]) -> Path:
    """Generate Markdown (.md) meeting minutes."""
    out_file = Path(output_path)
    out_file.parent.mkdir(parents=True, exist_ok=True)

    lines = [
        f"# {minutes.get('title', 'Meeting Minutes')}",
        "",
        f"**Date:** {minutes.get('date', datetime.now().strftime('%B %d, %Y'))}  ",
        f"**Duration:** {minutes.get('duration_formatted', 'N/A')}  ",
        f"**Attendees:** {', '.join(minutes.get('participants', ['Unknown']))}",
        "",
        "---",
        "",
        "## Executive Summary",
        "",
        minutes.get("summary", "No summary available."),
        "",
        "## Key Discussion Points",
        ""
    ]

    for pt in minutes.get("key_points", []):
        lines.append(f"- {pt}")

    lines.extend(["", "## Decisions Made", ""])
    for dec in minutes.get("decisions", []):
        lines.append(f"- **{dec}**")

    lines.extend([
        "",
        "## Action Items",
        "",
        "| Task | Responsible Person | Deadline | Priority |",
        "| :--- | :--- | :--- | :--- |"
    ])
    for item in minutes.get("action_items", []):
        lines.append(
            f"| {item.get('task')} | {item.get('responsible_person')} | {item.get('deadline')} | {item.get('priority', 'Normal')} |"
        )

    lines.extend(["", "## Speaker-Wise Transcript", ""])
    for entry in minutes.get("speaker_contributions", []):
        lines.append(f"- **{entry.get('speaker')}**: {entry.get('text')}")

    out_file.write_text("\n".join(lines), encoding="utf-8")
    return out_file


def export_txt(minutes: Dict[str, Any], output_path: Union[str, Path]) -> Path:
    """Generate clean Plain Text (.txt) meeting minutes."""
    out_file = Path(output_path)
    out_file.parent.mkdir(parents=True, exist_ok=True)

    lines = [
        "=" * 60,
        minutes.get("title", "MEETING MINUTES").upper(),
        "=" * 60,
        f"Date     : {minutes.get('date', datetime.now().strftime('%B %d, %Y'))}",
        f"Duration : {minutes.get('duration_formatted', 'N/A')}",
        f"Attendees: {', '.join(minutes.get('participants', ['Unknown']))}",
        "=" * 60,
        "",
        "EXECUTIVE SUMMARY:",
        minutes.get("summary", ""),
        "",
        "KEY DISCUSSION POINTS:"
    ]

    for pt in minutes.get("key_points", []):
        lines.append(f" - {pt}")

    lines.extend(["", "DECISIONS MADE:"])
    for dec in minutes.get("decisions", []):
        lines.append(f" - [DECISION] {dec}")

    lines.extend(["", "ACTION ITEMS:"])
    for item in minutes.get("action_items", []):
        lines.append(
            f" - [TASK] {item.get('task')}\n"
            f"   Assignee: {item.get('responsible_person')} | Deadline: {item.get('deadline')} | Priority: {item.get('priority', 'Normal')}"
        )

    lines.extend(["", "TRANSCRIPT:"])
    for entry in minutes.get("speaker_contributions", []):
        lines.append(f"[{entry.get('speaker')}]: {entry.get('text')}")

    out_file.write_text("\n".join(lines), encoding="utf-8")
    return out_file


def export_all_formats(
    minutes: Dict[str, Any],
    output_dir: Union[str, Path],
    meeting_stem: str
) -> Dict[str, Path]:
    """Export meeting minutes into all formats: PDF, DOCX, MD, TXT, and JSON."""
    out_dir = Path(output_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    pdf_path = out_dir / f"{meeting_stem}_minutes.pdf"
    docx_path = out_dir / f"{meeting_stem}_minutes.docx"
    md_path = out_dir / f"{meeting_stem}_minutes.md"
    txt_path = out_dir / f"{meeting_stem}_minutes.txt"
    json_path = out_dir / f"{meeting_stem}_minutes.json"

    export_pdf(minutes, pdf_path)
    export_docx(minutes, docx_path)
    export_markdown(minutes, md_path)
    export_txt(minutes, txt_path)

    with json_path.open("w", encoding="utf-8") as f:
        json.dump(minutes, f, indent=2, ensure_ascii=False)

    return {
        "pdf": pdf_path,
        "docx": docx_path,
        "md": md_path,
        "txt": txt_path,
        "json": json_path
    }
